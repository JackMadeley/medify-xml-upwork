import os

from docx import Document
from docx.table import Table

from Models.Questions.SJQuestion import SJQuestion
from Utilities.Documents.Cell import Cell
from Utilities.Documents.Docx import Docx


class SJ_Docx(Docx):

    def __init__(self):
        super().__init__()
        self.template = os.path.join(self.template_folder, "SJ", "UCAT SJ Template.docx")

    def generate_template(self, question: SJQuestion):
        if type(question) != SJQuestion:
            raise Exception(f"A SJQuestion was expected but a {type(question)} was provided")
        document = Document(self.template)
        document_table: Table = document.tables[0]
        # Set the item name
        document_table.cell(0, 2).text = question.item_name
        # Set the category name
        document_table.cell(0, 5).text = question.category
        # Add stimulus
        stimulus = Cell(document_table.cell(2, 1))
        stimulus.add_contents_to_cells(question.stimulus)
        # Add question set members
        full_table = Cell(document_table)
        full_table.insert_question_set_members_4_answer(3, question.question_set_members)
        return document
